"""Write demo/demo.docx: a short fake paper with planted problems, for trying the checker.
No manuscript text. Run:  python demo/make_demo_docx.py
"""
from pathlib import Path

import docx

HERE = Path(__file__).parent


def main() -> Path:
    d = docx.Document()
    d.add_paragraph("Rules Accumulate Faster Than They Disappear in Online Communities", style="Title")
    d.add_paragraph("A. Author, Example University")

    d.add_heading("Abstract", 1)
    d.add_paragraph("Online communities utilize written rules in order to govern participation. "
                    "We demonstrate that rules evolve over time. This paper integrates two theories of change.")
    d.add_paragraph("Keywords governance, rules, online communities")

    d.add_heading("Introduction", 1)
    d.add_paragraph("Rules matter for communities. It is important to note that prior to 2020, "
                    "a number of communities had no written rules. The majority of communities added rules "
                    "(Smith, 2019; Jones et al., 2020a).")
    d.add_paragraph("Jones et al. (2020a) argue that modification of rules is rare, but revision and "
                    "amendment are common. Communities that changed their rules mostly added new ones over time.")
    d.add_paragraph("H1: Communities with more subscribers have no systematic difference in rule changes.")

    d.add_heading("Methods", 1)
    d.add_heading("Data", 2)
    d.add_paragraph("First, we conducted an analysis of 130,851 communities. There was a reduction in deletions. "
                    "The reduction in deletions was due to the decline in activity.")
    d.add_paragraph("Rules were coded by two raters. Rules were compared across snapshots. "
                    "Rules were then aggregated, as noted above.")
    d.add_paragraph("Table 1 Descriptive Statistics")
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(("Community", "n", "%")):
        t.cell(0, i).text = v
    d.add_paragraph("Note. N = 130,851 communities.")

    d.add_heading("Results", 1)
    d.add_paragraph("The bootstrapping analysis revealed a crucial pattern. Most communities primarily "
                    "added rules. This pattern was significant.")
    d.add_paragraph("Figure 1. Most rule changes are additions.")
    d.add_paragraph("A Wilcoxon test confirmed the difference (p < .001). Researchers were able to "
                    "ascertain the end result.")

    d.add_heading("Discussion", 1)
    d.add_paragraph("This shows that rules evolve. The vast majority of communities changed, and "
                    "communities that changed their rules mostly added new ones over time. In conclusion, "
                    "the results seem to suggest that Institutional Layering Theory applies.")

    d.add_heading("References", 1)
    d.add_paragraph("Jones, A., Brown, B., & Lee, C. (2020a). Rules in online communities. "
                    "Journal of Rules, 12(3), 100–120. https://doi.org/10.1000/example.2020")
    d.add_paragraph("Smith, D. (2019). A study of governance. Governance Review, 4(1), 1–20.")
    out = HERE / "demo.docx"
    d.save(out)
    return out


if __name__ == "__main__":
    print(main())
